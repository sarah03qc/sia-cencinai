import json
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"

FINAL_XLSX = DATA / "CENCINAI_Benchmark_Final_299.xlsx"
EVIDENCE_XLSX = DATA / "Evaluacion_Retrieval_CENCINAI_300.xlsx"

OUT_BENCHMARK = DATA / "benchmark_final_299.json"
OUT_EVIDENCE = DATA / "evidence_annotations_final_299.json"

EXCLUDED_ID = "sino_023"


def sheet_records(ws):
    rows = ws.iter_rows(values_only=True)

    headers = next(rows)
    headers = [
        str(x).strip() if x is not None else None
        for x in headers
    ]

    records = []

    for row in rows:
        if not any(v is not None for v in row):
            continue

        record = {}

        for header, value in zip(headers, row):
            if header is not None:
                record[header] = value

        records.append(record)

    return records


def clean(value):
    if value is None:
        return ""
    return str(value).strip()


# ============================================================
# 1 BENCHMARK FINAL 299
# ============================================================

wb_final = load_workbook(
    FINAL_XLSX,
    read_only=True,
    data_only=True
)

ws_final = wb_final["Benchmark_Final"]
final_rows = sheet_records(ws_final)

benchmark = []

for r in final_rows:
    benchmark.append({
        "id": clean(r["id"]),
        "categoria": clean(r["categoria"]),
        "pregunta": clean(r["pregunta"]),
        "ground_truth": clean(r["ground_truth"]),
    })


# ============================================================
# 2 VALIDACIONES DEL BENCHMARK
# ============================================================

if len(benchmark) != 299:
    raise ValueError(
        f"Esperaba 299 preguntas finales, encontré {len(benchmark)}"
    )

ids = [x["id"] for x in benchmark]

if len(ids) != len(set(ids)):
    raise ValueError("Hay IDs duplicados en Benchmark_Final")

if EXCLUDED_ID in ids:
    raise ValueError(
        f"{EXCLUDED_ID} todavía está presente en Benchmark_Final"
    )

benchmark_by_id = {
    x["id"]: x
    for x in benchmark
}


# ============================================================
# 3 EVIDENCIA DEL AUDIT ORIGINAL
# ============================================================

wb_ev = load_workbook(
    EVIDENCE_XLSX,
    read_only=True,
    data_only=True
)

ws_ev = wb_ev["Referencia_Evidencia"]
evidence_rows = sheet_records(ws_ev)

evidence_original_by_id = {}

for r in evidence_rows:
    qid = clean(r["ID"])

    if qid:
        evidence_original_by_id[qid] = r


# ============================================================
# 4 CONSTRUIR EVIDENCIA FINAL PARA SOLO LAS 299
# ============================================================

evidence_final = []
missing = []

for item in benchmark:

    qid = item["id"]

    if qid not in evidence_original_by_id:
        missing.append(qid)
        continue

    r = evidence_original_by_id[qid]

    evidence_final.append({
        "id": qid,
        "categoria": item["categoria"],
        "pregunta": item["pregunta"],

        # IMPORTANTE:
        # usar el GT FINAL auditado, no el GT viejo del Excel de 300
        "ground_truth": item["ground_truth"],

        "supported_by_corpus":
            clean(r.get("Supported by corpus")),

        "answer_from_corpus":
            clean(r.get("Answer from corpus")),

        "source_document":
            clean(r.get("Documento(s) fuente")),

        "page_or_section":
            clean(r.get("Página / sección")),

        "supporting_excerpt":
            clean(r.get("Evidencia de soporte")),

        "notebooklm_notes":
            clean(r.get("Notas NotebookLM")),

        "binary_corpus_answer":
            clean(r.get("Respuesta binaria corpus")),

        "original_gt_vs_corpus":
            clean(r.get("GT vs corpus")),

        "original_review_reason":
            clean(r.get("Motivo de revisión")),
    })


if missing:
    raise ValueError(
        f"Faltan evidencias para {len(missing)} IDs: {missing[:20]}"
    )

if len(evidence_final) != 299:
    raise ValueError(
        f"Esperaba 299 evidencias finales, obtuve {len(evidence_final)}"
    )


# ============================================================
# 5 VALIDACIONES IMPORTANTES
# ============================================================

categories = {}

for x in benchmark:
    categories[x["categoria"]] = (
        categories.get(x["categoria"], 0) + 1
    )


support_counts = {}

for x in evidence_final:
    status = x["supported_by_corpus"]
    support_counts[status] = (
        support_counts.get(status, 0) + 1
    )


without_document = [
    x["id"]
    for x in evidence_final
    if not x["source_document"]
]

without_evidence = [
    x["id"]
    for x in evidence_final
    if not x["supporting_excerpt"]
]


# Comprobar las preguntas sí/no:
# el GT final debería coincidir con la respuesta binaria del corpus
binary_mismatches = []

for x in evidence_final:

    if x["categoria"] != "sino":
        continue

    corpus_answer = (
        x["binary_corpus_answer"]
        .strip()
        .lower()
    )

    final_gt = (
        x["ground_truth"]
        .strip()
        .lower()
    )

    # normalización mínima
    corpus_answer = (
        corpus_answer
        .replace("si", "sí")
        if corpus_answer == "si"
        else corpus_answer
    )

    final_gt = (
        final_gt
        .replace("si", "sí")
        if final_gt == "si"
        else final_gt
    )

    if corpus_answer and corpus_answer != final_gt:
        binary_mismatches.append({
            "id": x["id"],
            "final_gt": x["ground_truth"],
            "corpus_answer": x["binary_corpus_answer"]
        })


# ============================================================
# 6 GUARDAR
# ============================================================

with open(
    OUT_BENCHMARK,
    "w",
    encoding="utf-8"
) as f:
    json.dump(
        benchmark,
        f,
        ensure_ascii=False,
        indent=2
    )


with open(
    OUT_EVIDENCE,
    "w",
    encoding="utf-8"
) as f:
    json.dump(
        evidence_final,
        f,
        ensure_ascii=False,
        indent=2
    )


# ============================================================
# 7 REPORTE
# ============================================================

print()
print("=" * 65)
print("BENCHMARK FINAL")
print("=" * 65)

print(f"Preguntas: {len(benchmark)}")

for cat, count in sorted(categories.items()):
    print(f"  {cat}: {count}")


print()
print("=" * 65)
print("EVIDENCIA")
print("=" * 65)

print(f"Anotaciones: {len(evidence_final)}")

for status, count in support_counts.items():
    print(
        f"  Supported by corpus {repr(status)}: {count}"
    )

print(f"Sin documento fuente: {len(without_document)}")
print(f"Sin evidencia textual: {len(without_evidence)}")

print(
    "Mismatch GT final vs respuesta binaria corpus: "
    f"{len(binary_mismatches)}"
)

if binary_mismatches:
    print("\nMismatches:")
    for x in binary_mismatches:
        print(x)


print()
print("=" * 65)
print("ARCHIVOS CREADOS")
print("=" * 65)

print(OUT_BENCHMARK)
print(OUT_EVIDENCE)
